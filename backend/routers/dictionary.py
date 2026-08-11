import io
import os
import re
import base64
import asyncio
import requests
import openpyxl
from fastapi import APIRouter, HTTPException, Depends, Request
from auth import get_user_token
from pydantic import BaseModel
from typing import List, Dict, Any

router = APIRouter()

DICT_VOLUME_PATH = '/Volumes/catalog_central/dictionary/file_upload'

# Sheet routing
DICT_SHEETS = ['System Level', 'Domain Level - All']
CODE_SHEETS  = ['System - Code Table', 'Domain - Code Table']


def _get_service_principal_token() -> str:
    token = os.getenv("DATABRICKS_TOKEN", "")
    if token:
        return token

    host = os.getenv("DATABRICKS_HOST", "").rstrip("/")
    client_id = os.getenv("DATABRICKS_CLIENT_ID", "")
    client_secret = os.getenv("DATABRICKS_CLIENT_SECRET", "")

    if not host or not client_id or not client_secret:
        return ""

    resp = requests.post(
        f"{host}/oidc/v1/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "all-apis",
        },
        timeout=10,
    )
    if resp.ok:
        return resp.json().get("access_token", "")
    return ""


def _get_db_config(request: Request):
    host  = os.getenv("DATABRICKS_HOST", "").rstrip("/")
    if not host:
        raise HTTPException(
            500,
            "DATABRICKS_HOST missing from environment."
        )

    service_token = _get_service_principal_token()
    if not service_token:
        raise HTTPException(
            500,
            "No service principal token available. "
            "Set DATABRICKS_TOKEN or DATABRICKS_CLIENT_ID and DATABRICKS_CLIENT_SECRET."
        )

    return host, service_token


def normalize_element_name(name: str) -> str:
    name = name.strip()
    name = re.sub(r"[ \-/\\]+", "_", name)
    name = re.sub(r"[\'\(\)\[\]\{\}\.\,\;\:\!\?\@\#\$\%\^\&\*\+\=\~\`\"]+", "", name)
    name = re.sub(r"[^a-zA-Z0-9_]", "", name)
    name = re.sub(r"_+", "_", name)
    name = name.strip("_")
    if name and name[0].isdigit():
        name = f"col_{name}"
    return name or "unnamed"


def _parse_sheet(ws) -> tuple[list, list]:
    """Return (headers, rows) from a worksheet, skipping fully-empty rows."""
    header_row  = next(ws.iter_rows(min_row=1, max_row=1))
    headers = [cell.value for cell in header_row if cell.value is not None]

    rows: list[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        obj: Dict[str, Any] = {}
        for col_idx, header in enumerate(headers):
            val = row[col_idx] if col_idx < len(row) else None
            str_val = str(val).strip() if val is not None else ''
            obj[header] = '' if str_val in ('None', '\xa0') else str_val
        rows.append(obj)
    return headers, rows


def _merge_sheets(wb, sheet_names: list[str]) -> tuple[list, list]:
    """Merge multiple sheets that share the same headers into one list."""
    combined_headers: list | None = None
    combined_rows: list = []

    for name in sheet_names:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        headers, rows = _parse_sheet(ws)
        if combined_headers is None:
            combined_headers = headers
        combined_rows.extend(rows)

    return combined_headers or [], combined_rows


class DictFileUpload(BaseModel):
    filename: str
    content_b64: str

@router.post("/parse-dictionary")
async def parse_dictionary(body: DictFileUpload):
    contents = base64.b64decode(body.content_b64)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)

        # Data Dictionary: System Level + Domain Level - All
        dd_headers, dd_rows = _merge_sheets(wb, DICT_SHEETS)
        for row in dd_rows:
            elem = row.get('Data Element Name', '')
            row['_safe_element_name'] = normalize_element_name(elem) if elem else ''
            row['_generated_sql']     = ''

        # Master Code: System - Code Table + Domain - Code Table
        mc_headers, mc_rows = _merge_sheets(wb, CODE_SHEETS)

        return {
            "data_dictionary": {"headers": dd_headers, "rows": dd_rows},
            "master_code":      {"headers": mc_headers, "rows": mc_rows},
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel: {str(e)}")


class SheetData(BaseModel):
    headers: List[str]
    rows: List[Dict[str, Any]]


class UploadDictionaryRequest(BaseModel):
    data_dictionary: SheetData
    master_code:     SheetData
    filename: str = "NCSS_Data_Dictionary.xlsx"


def _cell_value(row: dict, h: str) -> str:
    """Resolve a single cell value, applying normalized name / generated SQL overrides."""
    if h == 'Data Element Name':
        return row.get('_safe_element_name') or row.get(h, '')
    if h == 'Dictionary Technical Logic':
        sql = (row.get('_generated_sql') or '').strip()
        return sql if sql else row.get(h, '')
    return row.get(h, '')


def _write_sheet(wb, title: str, headers: list, rows: list, first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append([_cell_value(row, h) for h in headers])


def _upload_bytes(host: str, token: str, full_path: str, content: bytes) -> requests.Response:
    url = f"{host}/api/2.0/fs/files{full_path}"
    return requests.put(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params={"overwrite": "true"},
        data=content,
    )


def _build_and_upload_sync(
    host: str, token: str,
    dd: SheetData, mc: SheetData,
    filename: str,
) -> str:
    wb = openpyxl.Workbook()
    dd_pub = [h for h in dd.headers if not h.startswith('_')]

    # Split DD rows back into original sheets using Dictionary Level discriminator
    system_rows = [r for r in dd.rows if r.get('Dictionary Level', '') == 'SYSTEM']
    domain_rows  = [r for r in dd.rows if r.get('Dictionary Level', '') != 'SYSTEM']

    _write_sheet(wb, 'System Level',      dd_pub, system_rows, first=True)
    _write_sheet(wb, 'Domain Level - All', dd_pub, domain_rows)

    # Split MC rows back into original sheets using Code Level discriminator
    sys_code_rows = [r for r in mc.rows if r.get('Code Level', '') == 'System']
    dom_code_rows  = [r for r in mc.rows if r.get('Code Level', '') != 'System']

    _write_sheet(wb, 'System - Code Table', mc.headers, sys_code_rows)
    _write_sheet(wb, 'Domain - Code Table', mc.headers, dom_code_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    full_path = f"{DICT_VOLUME_PATH}/{filename}"
    content = buf.getvalue()

    resp = _upload_bytes(host, token, full_path, content)
    if resp.status_code in (200, 201, 204):
        return full_path
    raise HTTPException(
        status_code=resp.status_code,
        detail=f"Databricks upload failed ({resp.status_code}): {resp.text}",
    )


@router.post("/upload-dictionary")
async def upload_dictionary(req: UploadDictionaryRequest, request: Request, _: str = Depends(get_user_token)):
    host, service_token = _get_db_config(request)
    try:
        full_path = await asyncio.to_thread(
            _build_and_upload_sync, host, service_token,
            req.data_dictionary, req.master_code, req.filename
        )
        return {"status": "uploaded", "path": full_path}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
